#!/usr/bin/env python3
"""
Build daily, versioned CSV and Excel files containing ECB fiat reference rates and
CoinMarketCap cryptocurrency rates.

Rate convention (important):
    Every rate column is "units of that asset/currency per 1 EUR".

Examples:
    USD = 1.1205    means 1 EUR = 1.1205 USD
    BTC = 0.000014  means 1 EUR = 0.000014 BTC

Calendar behavior:
- ECB publishes fiat rates around 16:00 Central European local time on banking working days.
- The output includes EVERY calendar day through yesterday (Europe/Amsterdam).
- Non-ECB days (weekends/holidays) have blank fiat cells.
- Crypto cells use completed daily prices where data is available, including weekends.
- Dates before a coin's CoinMarketCap first_historical_data remain blank.
- If a crypto download is blocked after 3 attempts, affected cells are "ERROR".

Recovery behavior:
- ECB ZIPs are cached and refreshed when a newer scheduled release is expected,
  so missed GitHub Actions windows are repaired by the full historical file.
- The latest prior output snapshot is used as a crypto cache.
- Missing or ERROR crypto cells are retried, including older gaps.
- Filenames use the run date; data ends on the previous calendar day.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import os
import re
import sys
import time
import zipfile
from dataclasses import dataclass
from decimal import Decimal
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from zoneinfo import ZoneInfo

import pandas as pd
import requests

ECB_DAILY_ZIP = "https://www.ecb.europa.eu/stats/eurofxref/eurofxref.zip"
ECB_HIST_ZIP = "https://www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.zip"

CMC_MAP_PUBLIC = "https://pro-api.coinmarketcap.com/public-api/v1/cryptocurrency/map"
CMC_MAP_PRO = "https://pro-api.coinmarketcap.com/v1/cryptocurrency/map"
CMC_PRO_HIST = "https://pro-api.coinmarketcap.com/v3/cryptocurrency/quotes/historical"

# Unofficial fallback used by coinmarketcap.com itself. It can change or be
# bot-blocked. Prefer the official API by setting CMC_API_KEY in GitHub Secrets.
CMC_WEB_HIST = "https://api.coinmarketcap.com/data-api/v3/cryptocurrency/historical"

EUR_CMC_ID = 2790
USD_CMC_ID = 2781
OUTPUT_PREFIX = "eurofxref_and_crypto_rates"
OUTPUT_DIR = Path("output")
ASSET_CONFIG = Path("assets.json")
ECB_CACHE_DIR = Path(".cache/ecb")
LOCAL_TZ = ZoneInfo("Europe/Amsterdam")

# downloader settings
MAX_ATTEMPTS = 3
CONNECT_TIMEOUT = 10 # seconds
READ_TIMEOUT = 45
REQUEST_PAUSE_SECONDS = 0.35    # interval
CHUNK_DAYS = 180

# sunshine is the best medicine
USER_AGENT = (
    "crypto-exchangerate-script/1.0 "
    "(https://github.com/ikolar/crypto-exchangerate-script; daily reference-data job)"
)

# keywords used to check whether we were blocked when attempting to download
# you never know ...
BLOCK_HINTS = (
    "captcha",
    "cloudflare",
    "access denied",
    "temporarily blocked",
    "unusual traffic",
    "rate limit",
    "too many requests",
)

class BlockedDownload(RuntimeError):
    pass


class DownloadFailed(RuntimeError):
    pass


@dataclass(frozen=True)
class Asset:
    symbol: str
    name: str
    cmc_id: int
    first_historical_date: pd.Timestamp | None


def log(msg: str) -> None:
    print(msg, flush=True)


def today_local() -> pd.Timestamp:
    return pd.Timestamp(datetime.now(LOCAL_TZ).date())


def make_session() -> requests.Session:
    s = requests.Session()
    s.headers.update(
        {
            "User-Agent": USER_AGENT,
            "Accept": "application/json,text/csv,application/zip,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.8",
            "Connection": "keep-alive",
        }
    )
    return s


def looks_blocked(resp: requests.Response) -> bool:
    if resp.status_code in (403, 429):
        return True
    ctype = (resp.headers.get("content-type") or "").lower()
    # JSON endpoints returning HTML are commonly a WAF/challenge page.
    if "json" in resp.request.headers.get("Accept", "").lower() and "text/html" in ctype:
        return True
    sample = (resp.text[:2000] if "text" in ctype or "html" in ctype else "").lower()
    return any(h in sample for h in BLOCK_HINTS)


def request_with_retries(
    session: requests.Session,
    url: str,
    *,
    params: dict | None = None,
    headers: dict | None = None,
    expect_json: bool = False,
) -> requests.Response:
    last_exc: Exception | None = None
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            hdrs = dict(headers or {})
            if expect_json:
                hdrs["Accept"] = "application/json"
            resp = session.get(
                url,
                params=params,
                headers=hdrs,
                timeout=(CONNECT_TIMEOUT, READ_TIMEOUT),
            )
            if looks_blocked(resp):
                last_exc = BlockedDownload(
                    f"blocked/rate-limited: HTTP {resp.status_code} from {url}"
                )
            elif resp.status_code >= 500:
                last_exc = DownloadFailed(f"HTTP {resp.status_code} from {url}")
            elif resp.status_code >= 400:
                # Retry ordinary 4xx too; APIs sometimes return transient 401/403/429
                # through gateways. After the final try this becomes a hard failure.
                last_exc = DownloadFailed(
                    f"HTTP {resp.status_code} from {url}: {resp.text[:250]}"
                )
            else:
                if expect_json:
                    try:
                        resp.json()
                    except Exception as exc:
                        last_exc = BlockedDownload(
                            f"expected JSON but got non-JSON response from {url}"
                        )
                    else:
                        return resp
                else:
                    return resp
        except requests.RequestException as exc:
            last_exc = exc

        if attempt < MAX_ATTEMPTS:
            delay = (2, 5)[attempt - 1]
            log(f"  attempt {attempt}/{MAX_ATTEMPTS} failed; retrying in {delay}s: {last_exc}")
            time.sleep(delay)

    if isinstance(last_exc, BlockedDownload):
        raise last_exc
    raise DownloadFailed(str(last_exc) if last_exc else f"download failed: {url}")


def read_csv_from_zip(content: bytes, expected_name_fragment: str) -> pd.DataFrame:
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        names = [n for n in zf.namelist() if n.lower().endswith(".csv")]
        if not names:
            raise DownloadFailed("ZIP contained no CSV")
        preferred = [n for n in names if expected_name_fragment.lower() in n.lower()]
        name = preferred[0] if preferred else names[0]
        with zf.open(name) as fh:
            return pd.read_csv(fh)


def clean_ecb_frame(df: pd.DataFrame, daily: bool) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df.loc[:, [c for c in df.columns if c and not c.lower().startswith("unnamed")]]

    if "Date" not in df.columns:
        raise DownloadFailed(f"ECB CSV missing Date column: {df.columns.tolist()}")

    if daily:
        # ECB daily CSV may use e.g. "22 October 2026".
        dates = pd.to_datetime(df["Date"], errors="coerce", dayfirst=True)
    else:
        dates = pd.to_datetime(df["Date"], format="%Y-%m-%d", errors="coerce")

    df["Date"] = dates
    df = df.dropna(subset=["Date"]).set_index("Date")

    for col in df.columns:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    return df.sort_index()


def latest_expected_ecb_release(now: datetime | None = None) -> pd.Timestamp:
    """Latest 16:00 Europe/Amsterdam publication, excluding TARGET closing days."""
    local = (now or datetime.now(LOCAL_TZ)).astimezone(LOCAL_TZ)
    day = pd.Timestamp(local.date())
    if local.hour < 16:
        day -= pd.Timedelta(1, unit="D")
    while True:
        easter = pd.Timestamp(year=day.year, month=1, day=1) + pd.offsets.Easter()
        holidays = {
            pd.Timestamp(year=day.year, month=m, day=d)
            for m, d in ((1, 1), (5, 1), (12, 25), (12, 26))
        }
        holidays.update((easter - pd.Timedelta(2, unit="D"), easter + pd.Timedelta(1, unit="D")))
        if day.weekday() < 5 and day not in holidays:
            return day
        day -= pd.Timedelta(1, unit="D")


def download_ecb(
    session: requests.Session,
    daily_file: Path | None,
    hist_file: Path | None,
    *,
    cache_dir: Path = ECB_CACHE_DIR,
    force: bool = False,
    now: datetime | None = None,
) -> pd.DataFrame:
    expected = latest_expected_ecb_release(now)
    downloads = 0

    def load_source(url: str, local_file: Path | None, daily: bool) -> pd.DataFrame:
        nonlocal downloads
        if local_file is not None:
            return clean_ecb_frame(pd.read_csv(local_file), daily=daily)

        name = url.rsplit("/", 1)[-1]
        cached = cache_dir / name
        # Also recognize downloads left in the working directory by earlier runs.
        candidates = list(dict.fromkeys((cached, Path(name), Path(name).with_suffix(".csv"))))
        if not force:
            for candidate in candidates:
                if not candidate.exists():
                    continue
                try:
                    raw = (
                        pd.read_csv(candidate) if candidate.suffix == ".csv"
                        else read_csv_from_zip(candidate.read_bytes(), "eurofxref")
                    )
                    frame = clean_ecb_frame(raw, daily=daily)
                    if not frame.empty and frame.index.max() >= expected:
                        log(f"Using cached ECB data: {candidate} (through {frame.index.max().date()})")
                        return frame
                except (OSError, ValueError, zipfile.BadZipFile, pd.errors.ParserError, DownloadFailed) as exc:
                    log(f"WARNING: cannot use ECB cache {candidate}: {exc}")

        log(f"Downloading ECB {'daily' if daily else 'historical'} ZIP (expected through {expected.date()})...")
        response = request_with_retries(session, url)
        frame = clean_ecb_frame(read_csv_from_zip(response.content, "eurofxref"), daily=daily)
        if frame.empty:
            raise DownloadFailed(f"ECB ZIP contained no valid dated rows: {url}")
        cache_dir.mkdir(parents=True, exist_ok=True)
        staged = cached.with_suffix(".zip.tmp")
        staged.write_bytes(response.content)
        staged.replace(cached)
        downloads += 1
        if frame.index.max() < expected:
            log(f"WARNING: ECB source is still behind its expected {expected.date()} release; the next run will retry")
        return frame

    hist = load_source(ECB_HIST_ZIP, hist_file, daily=False)
    daily = load_source(ECB_DAILY_ZIP, daily_file, daily=True)
    if not downloads:
        log("No ECB downloads needed. Use --force to force an ECB download (explicit local CSV overrides stay local).")

    # Union columns because the historical file contains legacy currencies.
    all_cols = list(dict.fromkeys(list(hist.columns) + list(daily.columns)))
    hist = hist.reindex(columns=all_cols)
    daily = daily.reindex(columns=all_cols)

    # Daily row wins if it is newer or has a corrected same-day value.
    combined = pd.concat([hist, daily]).sort_index()
    combined = combined[~combined.index.duplicated(keep="last")]
    return combined


def load_asset_config() -> list[dict]:
    cfg = json.loads(ASSET_CONFIG.read_text(encoding="utf-8"))
    return cfg["major_cryptocurrencies"] + cfg["stablecoins"]


def load_currency_metadata() -> list[dict]:
    cfg = json.loads(ASSET_CONFIG.read_text(encoding="utf-8"))
    return cfg.get("fiat", []) + cfg["major_cryptocurrencies"] + cfg["stablecoins"]


def resolve_assets(session: requests.Session, config: list[dict]) -> list[Asset]:
    log("Resolving CoinMarketCap asset IDs / first historical dates...")
    configured = {
        wanted["symbol"].upper(): Asset(
            wanted["symbol"].upper(), wanted["name"], int(wanted["cmc_id"]),
            pd.Timestamp(wanted["first_historical_date"]).normalize(),
        )
        for wanted in config
        if wanted.get("cmc_id") and wanted.get("first_historical_date")
    }
    pending = [w for w in config if w["symbol"].upper() not in configured]
    api_key = os.environ.get("CMC_API_KEY", "").strip()
    params = {
        "listing_status": "active,inactive",
        "aux": "first_historical_data,last_historical_data",
        "limit": 5000,
        "start": 1,
    }
    rows: list[dict] = []
    while pending:
        try:
            resp = request_with_retries(
                session, CMC_MAP_PRO if api_key else CMC_MAP_PUBLIC,
                params=params,
                headers={"X-CMC_PRO_API_KEY": api_key} if api_key else None,
                expect_json=True,
            )
            payload = resp.json()
            status = payload.get("status") or {}
            if status.get("error_code"):
                raise DownloadFailed(str(status.get("error_message") or status))
            page = payload.get("data") or []
        except (BlockedDownload, DownloadFailed) as exc:
            log(f"WARNING: CoinMarketCap map failed: {exc}")
            break
        rows.extend(page)
        matches = {
            (str(r.get("symbol", "")).upper(), str(r.get("name", "")).casefold())
            for r in rows
        }
        if all((w["symbol"].upper(), w["name"].casefold()) in matches for w in pending):
            break
        if len(page) < params["limit"]:
            break
        params["start"] += len(page)

    assets: list[Asset] = []
    for wanted in config:
        sym = wanted["symbol"].upper()
        expected_name = wanted["name"].casefold()
        if sym in configured:
            assets.append(configured[sym])
            continue

        candidates = [r for r in rows if str(r.get("symbol", "")).upper() == sym]
        if not candidates:
            log(f"WARNING: CMC map has no asset for {sym}; it will be ERROR where expected")
            assets.append(Asset(sym, wanted["name"], -1, None))
            continue

        exact = [r for r in candidates if str(r.get("name", "")).casefold() == expected_name]
        chosen = exact[0] if exact else sorted(
            candidates,
            key=lambda r: (0 if r.get("is_active") else 1, int(r.get("id", 10**12))),
        )[0]

        first_raw = chosen.get("first_historical_data")
        first = pd.to_datetime(first_raw, errors="coerce", utc=True)
        first_date = (
            pd.Timestamp(first.date())
            if first is not pd.NaT and not pd.isna(first)
            else None
        )
        assets.append(
            Asset(
                symbol=sym,
                name=str(chosen.get("name") or wanted["name"]),
                cmc_id=int(chosen["id"]),
                first_historical_date=first_date,
            )
        )
    return assets


def latest_existing_snapshot(output_dir: Path, run_date: pd.Timestamp) -> Path | None:
    pattern = re.compile(
        rf"^{re.escape(OUTPUT_PREFIX)}\.(\d{{4}})\.(\d{{2}})\.(\d{{2}})\.csv$"
    )
    found: list[tuple[pd.Timestamp, Path]] = []
    for path in output_dir.glob(f"{OUTPUT_PREFIX}.*.csv"):
        m = pattern.match(path.name)
        if not m:
            continue
        d = pd.Timestamp(f"{m.group(1)}-{m.group(2)}-{m.group(3)}")
        if d <= run_date:
            found.append((d, path))
    return max(found, default=(None, None), key=lambda x: x[0])[1]


def load_previous_crypto(path: Path | None, symbols: list[str]) -> pd.DataFrame:
    if not path or not path.exists():
        return pd.DataFrame(columns=symbols)
    log(f"Using prior snapshot as crypto cache: {path}")
    prev = pd.read_csv(path, dtype=str, keep_default_na=False)
    if "Date" not in prev.columns:
        return pd.DataFrame(columns=symbols)
    prev["Date"] = pd.to_datetime(prev["Date"], errors="coerce")
    prev = prev.dropna(subset=["Date"]).set_index("Date")
    out = pd.DataFrame(index=prev.index)
    for sym in symbols:
        out[sym] = prev[sym] if sym in prev.columns else ""
    return out


def contiguous_windows(dates: Iterable[pd.Timestamp]) -> list[tuple[pd.Timestamp, pd.Timestamp]]:
    ds = sorted(set(pd.Timestamp(d).normalize() for d in dates))
    if not ds:
        return []
    windows: list[tuple[pd.Timestamp, pd.Timestamp]] = []
    start = prev = ds[0]
    for d in ds[1:]:
        if d == prev + pd.Timedelta(1, unit="D"):
            prev = d
        else:
            windows.append((start, prev))
            start = prev = d
    windows.append((start, prev))
    return windows


def split_window(start: pd.Timestamp, end: pd.Timestamp, days: int = CHUNK_DAYS):
    cursor = start
    while cursor <= end:
        chunk_end = min(end, cursor + pd.Timedelta(days - 1, unit="D"))
        yield cursor, chunk_end
        cursor = chunk_end + pd.Timedelta(1, unit="D")


def _extract_cmc_web_quotes(payload: dict) -> list[tuple[pd.Timestamp, float]]:
    data = payload.get("data") or {}
    quotes = data.get("quotes") or []
    result = []
    for q in quotes:
        # The public site's historical endpoint has used timeOpen/timeClose and
        # quote.close. Keep parsing defensive because this endpoint is unofficial.
        stamp = q.get("timeClose") or q.get("timeOpen") or q.get("timestamp")
        quote = q.get("quote") or {}
        price = quote.get("close")
        if price is None:
            price = quote.get("price")
        if price is None:
            price = q.get("close")
        if price is None:
            price = q.get("price")
        if stamp is None or price is None:
            continue
        dt = pd.to_datetime(stamp, errors="coerce", utc=True)
        if pd.isna(dt):
            continue
        try:
            result.append((pd.Timestamp(dt.date()), float(price)))
        except (TypeError, ValueError):
            continue
    return result


def _extract_cmc_pro_quotes(
    payload: dict, cmc_id: int, currency: str = "EUR",
) -> list[tuple[pd.Timestamp, float]]:
    data = payload.get("data") or {}
    item = data.get(str(cmc_id)) or data.get(cmc_id) or {}
    quotes = item.get("quotes") or []
    result = []
    for q in quotes:
        stamp = q.get("timestamp")
        quote = q.get("quote") or {}
        converted = quote.get(currency) or {}
        price = converted.get("price")
        if stamp is None or price is None:
            continue
        dt = pd.to_datetime(stamp, errors="coerce", utc=True)
        if pd.isna(dt):
            continue
        result.append((pd.Timestamp(dt.date()), float(price)))
    return result


def fetch_crypto_window(
    session: requests.Session,
    asset: Asset,
    start: pd.Timestamp,
    end: pd.Timestamp,
    api_key: str | None,
    *,
    currency: str = "EUR",
) -> tuple[dict[pd.Timestamp, float], bool]:
    """
    Returns ({date: price per coin in currency}, blocked_or_failed).
    A failed/blocked window is deliberately marked ERROR by the caller.
    """
    if asset.cmc_id < 0:
        return {}, True
    if currency not in ("EUR", "USD"):
        raise ValueError(f"Unsupported quote currency: {currency}")

    try:
        if api_key:
            # Official API is preferred. Use a 24h interval anchored near EOD UTC.
            params = {
                "id": str(asset.cmc_id),
                "time_start": f"{start.date().isoformat()}T23:59:00Z",
                "time_end": f"{end.date().isoformat()}T23:59:00Z",
                "interval": "24h",
                "convert": currency,
                "skip_invalid": "true",
            }
            headers = {"X-CMC_PRO_API_KEY": api_key}
            resp = request_with_retries(
                session, CMC_PRO_HIST, params=params, headers=headers, expect_json=True
            )
        else:
            # Free fallback. This is intentionally conservative: one asset,
            # moderate chunks, serialized calls, descriptive UA, no concurrency.
            params = {
                "id": str(asset.cmc_id),
                "convertId": str(EUR_CMC_ID if currency == "EUR" else USD_CMC_ID),
                "timeStart": int(start.replace(tzinfo=timezone.utc).timestamp()),
                # Include the full final day.
                "timeEnd": int(
                    (end + pd.Timedelta(1, unit="D") - pd.Timedelta(1, unit="s"))
                    .replace(tzinfo=timezone.utc)
                    .timestamp()
                ),
            }
            resp = request_with_retries(
                session, CMC_WEB_HIST, params=params, expect_json=True
            )
        payload = resp.json()
        status = payload.get("status") or {}
        if str(status.get("error_code", 0)) != "0":
            raise DownloadFailed(str(status.get("error_message") or status))
        rows = (
            _extract_cmc_pro_quotes(payload, asset.cmc_id, currency)
            if api_key else _extract_cmc_web_quotes(payload)
        )

        return {
            d: price for d, price in rows
            if start <= d <= end and math.isfinite(price) and price > 0
        }, False
    except (BlockedDownload, DownloadFailed) as exc:
        log(f"WARNING: {asset.symbol} {currency} {start.date()}..{end.date()} failed: {exc}")
        return {}, True
    finally:
        time.sleep(REQUEST_PAUSE_SECONDS)


def build_crypto(
    session: requests.Session,
    calendar: pd.DatetimeIndex,
    previous: pd.DataFrame,
    assets: list[Asset],
    run_date: pd.Timestamp,
    usd_per_eur: pd.Series | None = None,
) -> pd.DataFrame:
    api_key = os.environ.get("CMC_API_KEY", "").strip() or None
    if api_key:
        log("Crypto source: official CoinMarketCap Pro API (CMC_API_KEY found).")
    else:
        log(
            "Crypto source: CoinMarketCap public website historical endpoint "
            "(no CMC_API_KEY; this fallback may be bot-blocked)."
        )

    crypto = pd.DataFrame("", index=calendar, columns=[a.symbol for a in assets], dtype=object)
    fx = (
        pd.to_numeric(usd_per_eur.reindex(calendar), errors="coerce")
        if usd_per_eur is not None else pd.Series(float("nan"), index=calendar)
    )

    # Preserve previously downloaded values and ERROR markers.
    if not previous.empty:
        common = crypto.index.intersection(previous.index)
        for sym in crypto.columns:
            if sym in previous.columns:
                crypto.loc[common, sym] = previous.loc[common, sym].astype(str)

    for asset in assets:
        first = asset.first_historical_date
        if first is not None:
            # Old snapshots may contain fallback ERRORs from before this coin existed.
            crypto.loc[calendar < first, asset.symbol] = ""
        if first is None:
            # We do not know when data should start. Only mark dates after the
            # earliest CMC listing era as ERROR if we cannot resolve the asset.
            target_dates = calendar[calendar >= pd.Timestamp("2013-04-28")]
        else:
            target_dates = calendar[(calendar >= first) & (calendar <= run_date)]

        # Missing and previous ERROR cells are retried. This repairs outages.
        needs = [
            d for d in target_dates
            if str(crypto.at[d, asset.symbol]).strip() in ("", "ERROR", "nan", "NaN")
        ]
        if not needs:
            continue

        if len(needs) > 5:
            missing_dates = ", ".join(d.strftime("%Y-%m-%d") for d in needs[:3])
            missing_dates += f", ..., {needs[-1].strftime('%Y-%m-%d')}"
        else:
            missing_dates = ", ".join(d.strftime("%Y-%m-%d") for d in needs)
        log(
            f"{asset.symbol}: {len(needs)} date(s) need download/repair: "
            f"{missing_dates}"
        )
        for w_start, w_end in contiguous_windows(needs):
            for c_start, c_end in split_window(w_start, w_end):
                prices, failed = fetch_crypto_window(
                    session, asset, c_start, c_end, api_key
                )
                if not failed:
                    # A successful retry with no observation is a blank, not
                    # a stale download failure (for example, today's open day).
                    stale_errors = (
                        (calendar >= c_start) & (calendar <= c_end)
                        & crypto[asset.symbol].eq("ERROR").to_numpy()
                    )
                    crypto.loc[stale_errors, asset.symbol] = ""

                # Convert CMC's "1 coin = EUR X" to the ECB convention
                # "1 EUR = X coins".
                for d, eur_per_coin in prices.items():
                    if d in crypto.index and eur_per_coin and eur_per_coin > 0:
                        crypto.at[d, asset.symbol] = f"{1.0 / eur_per_coin:.12g}"

                if failed:
                    affected = calendar[(calendar >= c_start) & (calendar <= c_end)]
                    for d in affected:
                        if first is not None and d < first:
                            continue
                        # Never overwrite a previously good cached value.
                        cur = str(crypto.at[d, asset.symbol]).strip()
                        if cur in ("", "ERROR", "nan", "NaN"):
                            crypto.at[d, asset.symbol] = "ERROR"

        # Retry only unresolved dates with a valid same-day ECB USD rate.
        # USD per EUR / USD per coin = coins per EUR.
        usd_needs = [
            d for d in needs
            if str(crypto.at[d, asset.symbol]).strip() in ("", "ERROR", "nan", "NaN")
            and pd.notna(fx.at[d]) and math.isfinite(fx.at[d]) and fx.at[d] > 0
        ]
        if usd_needs:
            log(f"{asset.symbol}: retrying {len(usd_needs)} missing EUR date(s) in USD with same-day ECB conversion")
        for w_start, w_end in contiguous_windows(usd_needs):
            for c_start, c_end in split_window(w_start, w_end):
                prices, failed = fetch_crypto_window(
                    session, asset, c_start, c_end, api_key, currency="USD",
                )
                for d in usd_needs:
                    if not c_start <= d <= c_end:
                        continue
                    usd_per_coin = prices.get(d)
                    if usd_per_coin and math.isfinite(usd_per_coin) and usd_per_coin > 0:
                        crypto.at[d, asset.symbol] = f"{fx.at[d] / usd_per_coin:.12g}"
                    elif failed:
                        crypto.at[d, asset.symbol] = "ERROR"
                    else:
                        crypto.at[d, asset.symbol] = ""

    return crypto


def decimal_rate(value: object) -> str:
    """Serialize rates without exponents or rounding existing decimal digits."""
    if pd.isna(value) or str(value).strip() in ("", "nan", "NaN"):
        return ""
    if str(value).strip() == "ERROR":
        return "ERROR"
    number = Decimal(str(value))
    return format(number, "f") if number.is_finite() else ""


def add_excel_currency_headers(sheet, metadata: list[dict]) -> None:
    """Add a linked, human-readable second header row, retaining the symbol row."""
    from openpyxl.styles import Alignment

    currencies = {asset["symbol"]: asset for asset in metadata}
    if sheet["A2"].value != "Calendar date":
        sheet.insert_rows(2)
    sheet["A2"] = "Calendar date"
    for column in range(2, sheet.max_column + 1):
        symbol = sheet.cell(1, column).value
        asset = currencies.get(symbol, {})
        cell = sheet.cell(2, column, asset.get("name", symbol))
        url = asset.get("ecb_rates_url") or asset.get("cmc_url") or asset.get("ecb_url")
        if url:
            cell.hyperlink = url
            cell.style = "Hyperlink"
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 32
    sheet.freeze_panes = "B3"
    sheet.auto_filter.ref = f"A2:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"


def write_output(
    fiat: pd.DataFrame,
    crypto: pd.DataFrame,
    run_date: pd.Timestamp,
    output_dir: Path,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)

    merged = fiat.join(crypto, how="left")
    merged = merged.loc[merged.index < run_date.normalize()]
    # Older snapshots can contain an empty trailing ECB column.
    merged = merged.loc[:, [
        c for c in merged.columns
        if str(c).strip() and not str(c).strip().lower().startswith("unnamed")
    ]]
    merged.index.name = "Date"
    merged = merged.sort_index(ascending=False)

    out = output_dir / (
        f"{OUTPUT_PREFIX}.{run_date.year:04d}.{run_date.month:02d}.{run_date.day:02d}.csv"
    )

    # Format both numeric fiat cells and cached crypto strings as plain decimals.
    csv = merged.apply(lambda column: column.map(decimal_rate))
    csv.to_csv(out, na_rep="", date_format="%Y-%m-%d")
    # Use real numbers and dates in Excel, preserving blank cells and ERROR text.
    excel = merged.copy()
    for col in excel.columns:
        excel[col] = excel[col].map(
            lambda value: float(value)
            if pd.notna(value) and str(value).strip() not in ("", "ERROR")
            else ("ERROR" if str(value).strip() == "ERROR" else None)
        )
    with pd.ExcelWriter(out.with_suffix(".xlsx"), engine="openpyxl", datetime_format="YYYY-MM-DD") as writer:
        excel.to_excel(writer, sheet_name="Rates", na_rep="")
        sheet = writer.sheets["Rates"]
        add_excel_currency_headers(sheet, load_currency_metadata())
        sheet.column_dimensions["A"].width = 13
        for cell in sheet["A"][2:]:
            cell.number_format = "YYYY-MM-DD"
        for column in sheet.iter_cols(min_col=2):
            sheet.column_dimensions[column[0].column_letter].width = 24
            for cell in column[2:]:
                if cell.data_type == "n":
                    cell.number_format = "0.##############################"
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--ecb-cache-dir", type=Path, default=ECB_CACHE_DIR)
    parser.add_argument("--force", action="store_true", help="Redownload ECB ZIPs regardless of cache freshness.")
    parser.add_argument(
        "--run-date",
        help="Override snapshot run date (YYYY-MM-DD); data ends the previous day.",
    )
    parser.add_argument(
        "--ecb-daily-file",
        type=Path,
        help="Use a local unzipped ECB daily CSV instead of downloading it.",
    )
    parser.add_argument(
        "--ecb-hist-file",
        type=Path,
        help="Use a local unzipped ECB historical CSV instead of downloading it.",
    )
    args = parser.parse_args()

    run_date = (
        pd.Timestamp(args.run_date).normalize() if args.run_date else today_local()
    )
    data_end = run_date - pd.Timedelta(1, unit="D")

    session = make_session()

    # Reuse ECB downloads until the release schedule indicates newer data is due.
    ecb = download_ecb(
        session, args.ecb_daily_file, args.ecb_hist_file,
        cache_dir=args.ecb_cache_dir, force=args.force,
    )
    ecb = ecb[ecb.index <= data_end]

    if ecb.empty:
        raise RuntimeError("ECB dataset is empty after filtering")

    # The calendar starts at the ECB historical start (1999-01-04) and includes
    # every completed calendar day. Today's fiat rates are included next run.
    calendar = pd.date_range(ecb.index.min(), data_end, freq="D")

    # Reindexing is what creates blank weekend/holiday fiat rows.
    fiat = ecb.reindex(calendar)

    config = load_asset_config()
    symbols = [x["symbol"].upper() for x in config]
    previous_path = latest_existing_snapshot(args.output_dir, run_date)
    previous = load_previous_crypto(previous_path, symbols)

    assets = resolve_assets(session, config)

    crypto = build_crypto(session, calendar, previous, assets, data_end, fiat.get("USD"))
    out = write_output(fiat, crypto, run_date, args.output_dir)

    log(f"Created: {out}")
    log(f"Created: {out.with_suffix('.xlsx')}")
    log(f"Data through: {data_end.date()} (completed days only)")
    log(f"Rows: {len(calendar):,}; fiat columns: {fiat.shape[1]}; crypto columns: {crypto.shape[1]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
