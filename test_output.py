"""Verify CSV/Excel parity, blank-column removal, and spreadsheet cell types."""

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook
import pandas as pd

from rates import Asset, main, write_output, add_excel_currency_headers, load_currency_metadata, load_asset_config


class OutputTests(unittest.TestCase):
    def test_matching_csv_and_excel(self):
        dates = pd.date_range("2026-10-22", periods=3)
        fiat = pd.DataFrame({
            "USD": [1.12, 1.13, float("nan")],
            "Unnamed: 42": [float("nan")] * 3,
            " ": [float("nan")] * 3,
        }, index=dates)
        crypto = pd.DataFrame({"BTC": ["0.00001", "ERROR", ""]}, index=dates)
        with TemporaryDirectory() as directory:
            path = write_output(fiat, crypto, dates[-1] + pd.Timedelta(1, unit="D"), Path(directory))
            csv = pd.read_csv(path, dtype=str, keep_default_na=False)
            self.assertEqual(csv.columns.tolist(), ["Date", "USD", "BTC"])
            self.assertEqual(csv["Date"].tolist(), ["2026-10-24", "2026-10-23", "2026-10-22"])
            workbook = load_workbook(path.with_suffix(".xlsx"))
            sheet = workbook["Rates"]
            rows = list(sheet.values)
            self.assertEqual(list(rows[0]), csv.columns.tolist())
            self.assertEqual(rows[1], ("Calendar date", "US dollar", "Bitcoin"))
            self.assertEqual(rows[2], (datetime(2026, 10, 24), None, None))
            self.assertEqual(rows[3], (datetime(2026, 10, 23), 1.13, "ERROR"))
            self.assertEqual(rows[4], (datetime(2026, 10, 22), 1.12, 0.00001))
            self.assertEqual(sheet.freeze_panes, "B3")
            self.assertEqual(sheet.auto_filter.ref, "A2:C5")
            self.assertEqual(sheet["A3"].number_format, "YYYY-MM-DD")
            self.assertEqual(sheet["C5"].number_format, "0.##############################")
            self.assertEqual(sheet["B4"].number_format, "0.##############################")
            self.assertTrue(sheet["B2"].hyperlink.target.endswith("eurofxref-graph-usd.en.html"))
            self.assertEqual(sheet["C2"].hyperlink.target, "https://coinmarketcap.com/currencies/bitcoin/")
            add_excel_currency_headers(sheet, load_currency_metadata())
            self.assertEqual(sheet.max_row, 5)  # Reapplying metadata must not duplicate the header.
            workbook.close()

    def test_export_drops_current_and_future_cached_rows(self):
        dates = pd.date_range("2026-10-23", periods=3)
        fiat = pd.DataFrame({"USD": [1.1, 1.2, 1.3]}, index=dates)
        crypto = pd.DataFrame({"BTC": ["0.1", "0.2", "0.3"]}, index=dates)
        with TemporaryDirectory() as directory:
            path = write_output(fiat, crypto, dates[1], Path(directory))
            csv = pd.read_csv(path)
            self.assertEqual(csv["Date"].tolist(), ["2026-10-23"])
            book = load_workbook(path.with_suffix(".xlsx"))
            self.assertEqual(book["Rates"].max_row, 3)
            book.close()

    def test_fiat_metadata_is_not_resolved_as_crypto(self):
        from rates import resolve_assets

        metadata = load_currency_metadata()
        fiat = [a for a in metadata if "ecb_url" in a]
        self.assertEqual(len(fiat), 41)
        self.assertIn("CYP", [a["symbol"] for a in fiat])
        crypto = load_asset_config()
        self.assertTrue(all("ecb_url" not in a for a in crypto))
        with patch("rates.request_with_retries") as request, patch("rates.log"):
            assets = resolve_assets(Mock(), crypto)
        self.assertEqual(len(assets), len(crypto))
        request.assert_not_called()

    def test_csv_never_uses_scientific_notation(self):
        dates = pd.date_range("2026-10-20", periods=4)
        fiat = pd.DataFrame({"USD": [1e-20, 1e20, float("nan"), 1.12]}, index=dates)
        crypto = pd.DataFrame({
            "BTC": ["1.38166606522e-05", "9.876543210987654321E-30", "ERROR", ""],
        }, index=dates)
        with TemporaryDirectory() as directory:
            path = write_output(fiat, crypto, pd.Timestamp("2026-10-24"), Path(directory))
            csv = pd.read_csv(path, dtype=str, keep_default_na=False).set_index("Date")
            self.assertEqual(csv.at["2026-10-20", "BTC"], "0.0000138166606522")
            self.assertEqual(csv.at["2026-10-21", "BTC"], "0." + "0" * 29 + "9876543210987654321")
            self.assertEqual(csv.at["2026-10-20", "USD"], "0." + "0" * 19 + "1")
            self.assertEqual(csv.at["2026-10-21", "USD"], "100000000000000000000")
            self.assertEqual(csv.at["2026-10-22", "BTC"], "ERROR")
            self.assertEqual(csv.at["2026-10-22", "USD"], "")
            self.assertEqual(csv.at["2026-10-23", "BTC"], "")

    def test_daily_run_uses_yesterday_and_keeps_run_date_in_filename(self):
        for run_date in ("2026-10-24", "2027-01-01", "2024-03-01"):
            with self.subTest(run_date=run_date), TemporaryDirectory() as directory:
                run = pd.Timestamp(run_date)
                dates = pd.date_range(run - pd.Timedelta(2, unit="D"), run)
                fiat = pd.DataFrame({"USD": [1.1, 1.2, 1.3]}, index=dates)
                previous = pd.DataFrame({"BTC": ["0.1", "0.2", "0.3"]}, index=dates)
                with patch("sys.argv", ["rates.py", "--run-date", run_date, "--output-dir", directory]), patch(
                    "rates.download_ecb", return_value=fiat,
                ), patch("rates.make_session", return_value=Mock()), patch(
                    "rates.load_asset_config", return_value=[{"symbol": "BTC", "name": "Bitcoin"}],
                ), patch("rates.load_previous_crypto", return_value=previous), patch(
                    "rates.resolve_assets", return_value=[Asset("BTC", "Bitcoin", 1, dates[0])],
                ), patch("rates.fetch_crypto_window") as fetch, patch("rates.log"):
                    self.assertEqual(main(), 0)
                fetch.assert_not_called()
                path = next(Path(directory).glob("*.csv"))
                self.assertTrue(path.name.endswith(run.strftime(".%Y.%m.%d.csv")))
                csv = pd.read_csv(path)
                self.assertEqual(csv["Date"].tolist(), [d.strftime("%Y-%m-%d") for d in dates[:2][::-1]])
                self.assertEqual(csv["BTC"].tolist(), [0.2, 0.1])


if __name__ == "__main__":
    unittest.main()
